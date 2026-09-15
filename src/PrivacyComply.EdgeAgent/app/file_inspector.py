import csv
import os
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.file_profiler import (
    build_classification_summary,
    build_column_profiles,
)

from app.finding_builder import (
    build_finding_summary,
    build_findings,
)

from app.evidence_generator import (
    generate_inspection_evidence,
)


SUPPORTED_EXTENSIONS = {
    ".xlsx",
    ".csv",
}

DEFAULT_MAX_FILE_SIZE_MB = 200
DEFAULT_MAX_COLUMNS = 500
DEFAULT_MAX_ROWS = 100000
DEFAULT_MAX_CELLS = 5000000

DEFAULT_MAX_XLSX_UNCOMPRESSED_MB = 500
DEFAULT_MAX_XLSX_COMPRESSION_RATIO = 200


def get_max_file_size_bytes() -> int:
    configured_value = os.getenv(
        "EDGE_MAX_FILE_SIZE_MB",
        str(DEFAULT_MAX_FILE_SIZE_MB),
    )

    try:
        max_file_size_mb = int(
            configured_value
        )
    except ValueError as exception:
        raise RuntimeError(
            "EDGE_MAX_FILE_SIZE_MB must "
            "be a positive integer."
        ) from exception

    if max_file_size_mb <= 0:
        raise RuntimeError(
            "EDGE_MAX_FILE_SIZE_MB must "
            "be a positive integer."
        )

    return (
        max_file_size_mb
        * 1024
        * 1024
    )


def get_max_columns() -> int:
    configured_value = os.getenv(
        "EDGE_MAX_COLUMNS",
        str(DEFAULT_MAX_COLUMNS),
    )

    try:
        max_columns = int(
            configured_value
        )
    except ValueError as exception:
        raise RuntimeError(
            "EDGE_MAX_COLUMNS must "
            "be a positive integer."
        ) from exception

    if max_columns <= 0:
        raise RuntimeError(
            "EDGE_MAX_COLUMNS must "
            "be a positive integer."
        )

    return max_columns


def get_max_rows() -> int:
    configured_value = os.getenv(
        "EDGE_MAX_ROWS",
        str(DEFAULT_MAX_ROWS),
    )

    try:
        max_rows = int(
            configured_value
        )
    except ValueError as exception:
        raise RuntimeError(
            "EDGE_MAX_ROWS must "
            "be a positive integer."
        ) from exception

    if max_rows <= 0:
        raise RuntimeError(
            "EDGE_MAX_ROWS must "
            "be a positive integer."
        )

    return max_rows


def get_max_cells() -> int:
    configured_value = os.getenv(
        "EDGE_MAX_CELLS",
        str(DEFAULT_MAX_CELLS),
    )

    try:
        max_cells = int(
            configured_value
        )
    except ValueError as exception:
        raise RuntimeError(
            "EDGE_MAX_CELLS must "
            "be a positive integer."
        ) from exception

    if max_cells <= 0:
        raise RuntimeError(
            "EDGE_MAX_CELLS must "
            "be a positive integer."
        )

    return max_cells


def get_max_xlsx_uncompressed_bytes() -> int:
    configured_value = os.getenv(
        "EDGE_MAX_XLSX_UNCOMPRESSED_MB",
        str(
            DEFAULT_MAX_XLSX_UNCOMPRESSED_MB
        ),
    )

    try:
        max_uncompressed_mb = int(
            configured_value
        )
    except ValueError as exception:
        raise RuntimeError(
            "EDGE_MAX_XLSX_UNCOMPRESSED_MB "
            "must be a positive integer."
        ) from exception

    if max_uncompressed_mb <= 0:
        raise RuntimeError(
            "EDGE_MAX_XLSX_UNCOMPRESSED_MB "
            "must be a positive integer."
        )

    return (
        max_uncompressed_mb
        * 1024
        * 1024
    )


def get_max_xlsx_compression_ratio() -> int:
    configured_value = os.getenv(
        "EDGE_MAX_XLSX_COMPRESSION_RATIO",
        str(
            DEFAULT_MAX_XLSX_COMPRESSION_RATIO
        ),
    )

    try:
        max_compression_ratio = int(
            configured_value
        )
    except ValueError as exception:
        raise RuntimeError(
            "EDGE_MAX_XLSX_COMPRESSION_RATIO "
            "must be a positive integer."
        ) from exception

    if max_compression_ratio <= 0:
        raise RuntimeError(
            "EDGE_MAX_XLSX_COMPRESSION_RATIO "
            "must be a positive integer."
        )

    return max_compression_ratio


def validate_column_count(
    columns: list[str],
) -> None:
    max_columns = get_max_columns()

    if len(columns) > max_columns:
        raise HTTPException(
            status_code=413,
            detail=(
                "The uploaded file exceeds "
                "the configured maximum "
                "number of columns."
            ),
        )


def validate_row_count(
    data_rows: list[list[Any]],
) -> None:
    max_rows = get_max_rows()

    if len(data_rows) > max_rows:
        raise HTTPException(
            status_code=413,
            detail=(
                "The uploaded file exceeds "
                "the configured maximum "
                "number of data rows."
            ),
        )


def validate_cell_count(
    columns: list[str],
    data_rows: list[list[Any]],
) -> None:
    max_cells = get_max_cells()

    cell_count = (
        len(columns)
        * len(data_rows)
    )

    if cell_count > max_cells:
        raise HTTPException(
            status_code=413,
            detail=(
                "The uploaded file exceeds "
                "the configured maximum "
                "number of data cells."
            ),
        )


def validate_xlsx_archive(
    file_content: bytes,
) -> None:
    max_uncompressed_bytes = (
        get_max_xlsx_uncompressed_bytes()
    )

    max_compression_ratio = (
        get_max_xlsx_compression_ratio()
    )

    try:
        with ZipFile(
            BytesIO(file_content),
            mode="r",
        ) as archive:
            archive_entries = (
                archive.infolist()
            )

            total_uncompressed_bytes = sum(
                entry.file_size
                for entry in archive_entries
            )

            if (
                total_uncompressed_bytes
                > max_uncompressed_bytes
            ):
                raise HTTPException(
                    status_code=413,
                    detail=(
                        "The XLSX archive exceeds "
                        "the configured maximum "
                        "uncompressed size."
                    ),
                )

            for entry in archive_entries:
                if entry.file_size == 0:
                    continue

                if entry.compress_size == 0:
                    raise HTTPException(
                        status_code=413,
                        detail=(
                            "The XLSX archive contains "
                            "a suspiciously compressed "
                            "entry."
                        ),
                    )

                compression_ratio = (
                    entry.file_size
                    / entry.compress_size
                )

                if (
                    compression_ratio
                    > max_compression_ratio
                ):
                    raise HTTPException(
                        status_code=413,
                        detail=(
                            "The XLSX archive contains "
                            "a suspiciously compressed "
                            "entry."
                        ),
                    )

    except BadZipFile as exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "The XLSX file is invalid "
                "or corrupted."
            ),
        ) from exception


async def inspect_uploaded_file(
    file: UploadFile,
) -> dict[str, Any]:
    file_name = file.filename or ""

    file_extension = (
        Path(file_name)
        .suffix
        .lower()
    )

    if file_extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. "
                "Only .xlsx and .csv files are supported."
            ),
        )

    max_file_size_bytes = (
        get_max_file_size_bytes()
    )

    file_content = await file.read(
        max_file_size_bytes + 1
    )

    if (
        len(file_content)
        > max_file_size_bytes
    ):
        raise HTTPException(
            status_code=413,
            detail=(
                "The uploaded file exceeds "
                "the configured maximum file size."
            ),
        )

    if (
        file_extension == ".xlsx"
        and not file_content.startswith(b"PK")
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "The uploaded file content does not "
                "match the .xlsx file type."
            ),
        )

    if file_extension == ".xlsx":
        validate_xlsx_archive(
            file_content
        )

        return inspect_xlsx(
            file=file,
            file_name=file_name,
            file_extension=file_extension,
            file_content=file_content,
        )

    return inspect_csv(
        file=file,
        file_name=file_name,
        file_extension=file_extension,
        file_content=file_content,
    )


def inspect_xlsx(
    file: UploadFile,
    file_name: str,
    file_extension: str,
    file_content: bytes,
) -> dict[str, Any]:
    try:
        workbook = load_workbook(
            filename=BytesIO(file_content),
            read_only=True,
            data_only=True,
        )
    except (
        InvalidFileException,
        BadZipFile,
        OSError,
        ValueError,
        KeyError,
    ) as exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "The XLSX file is invalid "
                "or corrupted."
            ),
        ) from exception

    worksheet = workbook.active

    rows = [
        list(row)
        for row in worksheet.iter_rows(
            values_only=True,
        )
    ]

    if not rows:
        return build_empty_result(
            file=file,
            file_name=file_name,
            file_extension=file_extension,
            file_content=file_content,
            sheet_name=worksheet.title,
        )

    columns = [
        str(value).strip()
        if value is not None
        else ""
        for value in rows[0]
    ]

    validate_column_count(
        columns
    )

    data_rows = rows[1:]

    validate_row_count(
        data_rows
    )

    validate_cell_count(
        columns,
        data_rows,
    )

    return build_result(
        file=file,
        file_name=file_name,
        file_extension=file_extension,
        file_content=file_content,
        sheet_name=worksheet.title,
        columns=columns,
        data_rows=data_rows,
    )


def inspect_csv(
    file: UploadFile,
    file_name: str,
    file_extension: str,
    file_content: bytes,
) -> dict[str, Any]:
    try:
        decoded_content = file_content.decode(
            "utf-8-sig",
        )
    except UnicodeDecodeError as exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "The CSV file could not be decoded "
                "using UTF-8 encoding."
            ),
        ) from exception

    try:
        csv_reader = csv.reader(
            StringIO(decoded_content),
            strict=True,
        )

        rows = [
            list(row)
            for row in csv_reader
        ]

    except csv.Error as exception:
        raise HTTPException(
            status_code=400,
            detail=(
                "The CSV file is malformed "
                "and could not be parsed."
            ),
        ) from exception

    if not rows:
        return build_empty_result(
            file=file,
            file_name=file_name,
            file_extension=file_extension,
            file_content=file_content,
            sheet_name=None,
        )

    columns = [
        value.strip()
        for value in rows[0]
    ]

    validate_column_count(
        columns
    )

    data_rows = rows[1:]

    validate_row_count(
        data_rows
    )

    validate_cell_count(
        columns,
        data_rows,
    )

    return build_result(
        file=file,
        file_name=file_name,
        file_extension=file_extension,
        file_content=file_content,
        sheet_name=None,
        columns=columns,
        data_rows=data_rows,
    )


def build_result(
    file: UploadFile,
    file_name: str,
    file_extension: str,
    file_content: bytes,
    sheet_name: str | None,
    columns: list[str],
    data_rows: list[list[Any]],
) -> dict[str, Any]:
    column_profiles = build_column_profiles(
        columns,
        data_rows,
    )

    classification_summary = (
        build_classification_summary(
            column_profiles,
        )
    )

    findings = build_findings(
        column_profiles,
    )

    finding_summary = build_finding_summary(
        findings,
    )

    evidence = generate_inspection_evidence(
        file_name=file_name,
        sheet_name=sheet_name,
        row_count=len(data_rows),
        column_count=len(columns),
        column_profiles=column_profiles,
        classification_summary=classification_summary,
    )

    return {
        "fileName": file_name,
        "fileExtension": file_extension,
        "contentType": file.content_type,
        "fileSizeBytes": len(file_content),
        "processingLocation":
            "LOCAL_EDGE_AGENT",
        "sheetName": sheet_name,
        "rowCount": len(data_rows),
        "columnCount": len(columns),
        "columns": columns,
        "columnProfiles": column_profiles,
        "classificationSummary":
            classification_summary,
        "findingSummary":
            finding_summary,
        "findings": findings,
        "evidence": evidence,
    }


def build_empty_result(
    file: UploadFile,
    file_name: str,
    file_extension: str,
    file_content: bytes,
    sheet_name: str | None,
) -> dict[str, Any]:
    return {
        "fileName": file_name,
        "fileExtension": file_extension,
        "contentType": file.content_type,
        "fileSizeBytes": len(file_content),
        "processingLocation":
            "LOCAL_EDGE_AGENT",
        "sheetName": sheet_name,
        "rowCount": 0,
        "columnCount": 0,
        "columns": [],
        "columnProfiles": [],
        "classificationSummary": {
            "totalColumns": 0,
            "classifiedColumns": 0,
            "unclassifiedColumns": 0,
            "classificationCoveragePercentage": 0.0,
        },
        "findingSummary": {
            "totalFindings": 0,
            "infoCount": 0,
            "warningCount": 0,
            "errorCount": 0,
        },
        "findings": [],
        "evidence": [],
    }