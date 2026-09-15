import asyncio
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile
import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from starlette.datastructures import Headers

from app.file_inspector import (
    build_empty_result,
    get_max_cells,
    get_max_columns,
    get_max_file_size_bytes,
    get_max_rows,
    get_max_xlsx_compression_ratio,
    get_max_xlsx_uncompressed_bytes,
    inspect_csv,
    inspect_uploaded_file,
    inspect_xlsx,
)


def create_upload_file(
    filename: str,
    content: bytes,
    content_type: str,
) -> UploadFile:
    return UploadFile(
        filename=filename,
        file=BytesIO(content),
        headers=Headers(
            {
                "content-type":
                    content_type,
            }
        ),
    )


def test_csv_inspection() -> None:
    csv_content = (
        "email,phone\n"
        "user1@example.com,+447700900001\n"
        "user2@example.com,+447700900002\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="test.csv",
        content=csv_content,
        content_type="text/csv",
    )

    result = inspect_csv(
        file=upload_file,
        file_name="test.csv",
        file_extension=".csv",
        file_content=csv_content,
    )

    assert result["fileName"] == "test.csv"
    assert result["fileExtension"] == ".csv"

    assert (
        result["processingLocation"]
        == "LOCAL_EDGE_AGENT"
    )

    assert result["rowCount"] == 2
    assert result["columnCount"] == 2

    assert result["columns"] == [
        "email",
        "phone",
    ]

    assert (
        result["columnProfiles"][0][
            "classificationCode"
        ]
        == "EMAIL_ADDRESS"
    )

    assert (
        result["columnProfiles"][1][
            "classificationCode"
        ]
        == "PHONE_NUMBER"
    )


def test_csv_header_whitespace_trimmed() -> None:
    csv_content = (
        " email , phone \n"
        "user@example.com,+447700900001\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="headers.csv",
        content=csv_content,
        content_type="text/csv",
    )

    result = inspect_csv(
        file=upload_file,
        file_name="headers.csv",
        file_extension=".csv",
        file_content=csv_content,
    )

    assert result["columns"] == [
        "email",
        "phone",
    ]


def test_empty_csv() -> None:
    csv_content = b""

    upload_file = create_upload_file(
        filename="empty.csv",
        content=csv_content,
        content_type="text/csv",
    )

    result = inspect_csv(
        file=upload_file,
        file_name="empty.csv",
        file_extension=".csv",
        file_content=csv_content,
    )

    assert result["rowCount"] == 0
    assert result["columnCount"] == 0
    assert result["columns"] == []
    assert result["columnProfiles"] == []
    assert result["findings"] == []

    assert result["classificationSummary"] == {
        "totalColumns": 0,
        "classifiedColumns": 0,
        "unclassifiedColumns": 0,
        "classificationCoveragePercentage": 0.0,
    }

    assert result["findingSummary"] == {
        "totalFindings": 0,
        "infoCount": 0,
        "warningCount": 0,
        "errorCount": 0,
    }


def test_invalid_utf8_csv_rejected() -> None:
    invalid_content = b"\xff\xfe\xfd"

    upload_file = create_upload_file(
        filename="invalid.csv",
        content=invalid_content,
        content_type="text/csv",
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        inspect_csv(
            file=upload_file,
            file_name="invalid.csv",
            file_extension=".csv",
            file_content=invalid_content,
        )

    assert (
        exception_info.value.status_code
        == 400
    )

    assert (
        exception_info.value.detail
        == (
            "The CSV file could not be decoded "
            "using UTF-8 encoding."
        )
    )


def test_xlsx_inspection() -> None:
    workbook = Workbook()
    worksheet = workbook.active

    worksheet.append(
        [
            "email",
            "phone",
        ]
    )

    worksheet.append(
        [
            "user1@example.com",
            "+447700900001",
        ]
    )

    worksheet.append(
        [
            "user2@example.com",
            "+447700900002",
        ]
    )

    workbook_buffer = BytesIO()
    workbook.save(workbook_buffer)

    xlsx_content = (
        workbook_buffer.getvalue()
    )

    upload_file = create_upload_file(
        filename="test.xlsx",
        content=xlsx_content,
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    result = inspect_xlsx(
        file=upload_file,
        file_name="test.xlsx",
        file_extension=".xlsx",
        file_content=xlsx_content,
    )

    assert result["fileName"] == "test.xlsx"
    assert result["fileExtension"] == ".xlsx"

    assert (
        result["processingLocation"]
        == "LOCAL_EDGE_AGENT"
    )

    assert result["rowCount"] == 2
    assert result["columnCount"] == 2

    assert result["columns"] == [
        "email",
        "phone",
    ]

    assert (
        result["columnProfiles"][0][
            "classificationCode"
        ]
        == "EMAIL_ADDRESS"
    )

    assert (
        result["columnProfiles"][1][
            "classificationCode"
        ]
        == "PHONE_NUMBER"
    )


def test_xlsx_uses_active_sheet_only() -> None:
    workbook = Workbook()

    first_sheet = workbook.active
    first_sheet.title = "Customers"

    first_sheet.append(
        [
            "email",
        ]
    )

    first_sheet.append(
        [
            "user@example.com",
        ]
    )

    second_sheet = workbook.create_sheet(
        title="Archive"
    )

    second_sheet.append(
        [
            "phone",
        ]
    )

    second_sheet.append(
        [
            "+447700900001",
        ]
    )

    workbook.active = 0

    workbook_buffer = BytesIO()
    workbook.save(workbook_buffer)

    xlsx_content = (
        workbook_buffer.getvalue()
    )

    upload_file = create_upload_file(
        filename="multi_sheet.xlsx",
        content=xlsx_content,
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    result = inspect_xlsx(
        file=upload_file,
        file_name="multi_sheet.xlsx",
        file_extension=".xlsx",
        file_content=xlsx_content,
    )

    assert (
        result["sheetName"]
        == "Customers"
    )

    assert result["columns"] == [
        "email",
    ]

    assert result["rowCount"] == 1

    assert (
        result["columnProfiles"][0][
            "classificationCode"
        ]
        == "EMAIL_ADDRESS"
    )


def test_unsupported_extension_rejected() -> None:
    upload_file = create_upload_file(
        filename="test.txt",
        content=b"hello",
        content_type="text/plain",
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        asyncio.run(
            inspect_uploaded_file(
                upload_file
            )
        )

    assert (
        exception_info.value.status_code
        == 400
    )

    assert (
        exception_info.value.detail
        == (
            "Unsupported file type. "
            "Only .xlsx and .csv files are supported."
        )
    )


def test_uploaded_csv_uses_full_pipeline() -> None:
    csv_content = (
        "business_identifier\n"
        "27ABCDE1234F1Z0\n"
        "29PQRSX5678K1ZU\n"
        "07AAAAA1111A1ZY\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="gstin.csv",
        content=csv_content,
        content_type="text/csv",
    )

    result = asyncio.run(
        inspect_uploaded_file(
            upload_file
        )
    )

    assert result["rowCount"] == 3
    assert result["columnCount"] == 1

    profile = (
        result["columnProfiles"][0]
    )

    assert (
        profile["classificationStatus"]
        == "CLASSIFIED"
    )

    assert (
        profile["classificationCode"]
        == "INDIA_GSTIN"
    )

    assert (
        profile["classificationMethod"]
        == "VALUE_PATTERN_RULE"
    )

    assert (
        profile["matchPercentage"]
        == 100.0
    )


def test_build_empty_result() -> None:
    upload_file = create_upload_file(
        filename="empty.csv",
        content=b"",
        content_type="text/csv",
    )

    result = build_empty_result(
        file=upload_file,
        file_name="empty.csv",
        file_extension=".csv",
        file_content=b"",
        sheet_name=None,
    )

    assert result == {
        "fileName": "empty.csv",
        "fileExtension": ".csv",
        "contentType": "text/csv",
        "fileSizeBytes": 0,
        "processingLocation":
            "LOCAL_EDGE_AGENT",
        "sheetName": None,
        "rowCount": 0,
        "columnCount": 0,
        "columns": [],
        "columnProfiles": [],
        "classificationSummary": {
            "totalColumns": 0,
            "classifiedColumns": 0,
            "unclassifiedColumns": 0,
            "classificationCoveragePercentage":
                0.0,
        },
        "findingSummary": {
            "totalFindings": 0,
            "infoCount": 0,
            "warningCount": 0,
            "errorCount": 0,
        },
        "findings": [],
        "evidence": [],
    }


def test_default_max_file_size(
    monkeypatch,
) -> None:
    monkeypatch.delenv(
        "EDGE_MAX_FILE_SIZE_MB",
        raising=False,
    )

    assert (
        get_max_file_size_bytes()
        == 200 * 1024 * 1024
    )


def test_configured_max_file_size(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_FILE_SIZE_MB",
        "10",
    )

    assert (
        get_max_file_size_bytes()
        == 10 * 1024 * 1024
    )


def test_invalid_max_file_size_configuration(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_FILE_SIZE_MB",
        "invalid",
    )

    with pytest.raises(
        RuntimeError,
        match="positive integer",
    ):
        get_max_file_size_bytes()


def test_oversized_uploaded_file_rejected(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_FILE_SIZE_MB",
        "1",
    )

    oversized_content = (
        b"a" * (1024 * 1024 + 1)
    )

    upload_file = create_upload_file(
        filename="oversized.csv",
        content=oversized_content,
        content_type="text/csv",
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        asyncio.run(
            inspect_uploaded_file(
                upload_file
            )
        )

    assert (
        exception_info.value.status_code
        == 413
    )

    assert (
        exception_info.value.detail
        == (
            "The uploaded file exceeds "
            "the configured maximum file size."
        )
    )


def test_invalid_xlsx_rejected() -> None:
    invalid_content = (
        b"This is not a real XLSX file."
    )

    upload_file = create_upload_file(
        filename="invalid.xlsx",
        content=invalid_content,
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        inspect_xlsx(
            file=upload_file,
            file_name="invalid.xlsx",
            file_extension=".xlsx",
            file_content=invalid_content,
        )

    assert (
        exception_info.value.status_code
        == 400
    )

    assert (
        exception_info.value.detail
        == (
            "The XLSX file is invalid "
            "or corrupted."
        )
    )


def test_blank_column_name_generates_finding() -> None:
    csv_content = (
        "email,,phone\n"
        "user1@example.com,value,+447700900001\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="blank_header.csv",
        content=csv_content,
        content_type="text/csv",
    )

    result = asyncio.run(
        inspect_uploaded_file(
            upload_file
        )
    )

    finding_codes = {
        finding["findingCode"]
        for finding in result["findings"]
    }

    assert (
        "BLANK_COLUMN_NAME"
        in finding_codes
    )


def test_duplicate_column_name_generates_finding() -> None:
    csv_content = (
        "email,email\n"
        "user1@example.com,user2@example.com\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="duplicate_header.csv",
        content=csv_content,
        content_type="text/csv",
    )

    result = asyncio.run(
        inspect_uploaded_file(
            upload_file
        )
    )

    duplicate_findings = [
        finding
        for finding in result["findings"]
        if finding["findingCode"]
        == "DUPLICATE_COLUMN_NAME"
    ]

    assert len(
        duplicate_findings
    ) == 1


def test_malformed_csv_rejected() -> None:
    malformed_content = (
        "email,name\n"
        "\"user@example.com,John\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="malformed.csv",
        content=malformed_content,
        content_type="text/csv",
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        inspect_csv(
            file=upload_file,
            file_name="malformed.csv",
            file_extension=".csv",
            file_content=malformed_content,
        )

    assert (
        exception_info.value.status_code
        == 400
    )

    assert (
        exception_info.value.detail
        == (
            "The CSV file is malformed "
            "and could not be parsed."
        )
    )


def test_xlsx_content_signature_mismatch_rejected() -> None:
    invalid_content = (
        b"This is plain text, not an XLSX file."
    )

    upload_file = create_upload_file(
        filename="fake.xlsx",
        content=invalid_content,
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        asyncio.run(
            inspect_uploaded_file(
                upload_file
            )
        )

    assert (
        exception_info.value.status_code
        == 400
    )

    assert (
        exception_info.value.detail
        == (
            "The uploaded file content does not "
            "match the .xlsx file type."
        )
    )


def test_default_max_columns(
    monkeypatch,
) -> None:
    monkeypatch.delenv(
        "EDGE_MAX_COLUMNS",
        raising=False,
    )

    assert (
        get_max_columns()
        == 500
    )


def test_invalid_max_columns_configuration(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_COLUMNS",
        "invalid",
    )

    with pytest.raises(
        RuntimeError,
        match="positive integer",
    ):
        get_max_columns()


def test_uploaded_file_exceeding_column_limit_rejected(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_COLUMNS",
        "2",
    )

    csv_content = (
        "column1,column2,column3\n"
        "value1,value2,value3\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="too_many_columns.csv",
        content=csv_content,
        content_type="text/csv",
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        asyncio.run(
            inspect_uploaded_file(
                upload_file
            )
        )

    assert (
        exception_info.value.status_code
        == 413
    )

    assert (
        exception_info.value.detail
        == (
            "The uploaded file exceeds "
            "the configured maximum "
            "number of columns."
        )
    )


def test_default_max_rows(
    monkeypatch,
) -> None:
    monkeypatch.delenv(
        "EDGE_MAX_ROWS",
        raising=False,
    )

    assert (
        get_max_rows()
        == 100000
    )


def test_invalid_max_rows_configuration(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_ROWS",
        "invalid",
    )

    with pytest.raises(
        RuntimeError,
        match="positive integer",
    ):
        get_max_rows()


def test_uploaded_file_exceeding_row_limit_rejected(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_ROWS",
        "2",
    )

    csv_content = (
        "email\n"
        "user1@example.com\n"
        "user2@example.com\n"
        "user3@example.com\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="too_many_rows.csv",
        content=csv_content,
        content_type="text/csv",
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        asyncio.run(
            inspect_uploaded_file(
                upload_file
            )
        )

    assert (
        exception_info.value.status_code
        == 413
    )

    assert (
        exception_info.value.detail
        == (
            "The uploaded file exceeds "
            "the configured maximum "
            "number of data rows."
        )
    )


def test_default_max_cells(
    monkeypatch,
) -> None:
    monkeypatch.delenv(
        "EDGE_MAX_CELLS",
        raising=False,
    )

    assert (
        get_max_cells()
        == 5000000
    )


def test_invalid_max_cells_configuration(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_CELLS",
        "invalid",
    )

    with pytest.raises(
        RuntimeError,
        match="positive integer",
    ):
        get_max_cells()


def test_uploaded_file_exceeding_cell_limit_rejected(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_CELLS",
        "4",
    )

    csv_content = (
        "column1,column2\n"
        "value1,value2\n"
        "value3,value4\n"
        "value5,value6\n"
    ).encode("utf-8")

    upload_file = create_upload_file(
        filename="too_many_cells.csv",
        content=csv_content,
        content_type="text/csv",
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        asyncio.run(
            inspect_uploaded_file(
                upload_file
            )
        )

    assert (
        exception_info.value.status_code
        == 413
    )

    assert (
        exception_info.value.detail
        == (
            "The uploaded file exceeds "
            "the configured maximum "
            "number of data cells."
        )
    )

def test_default_max_xlsx_uncompressed_size(
    monkeypatch,
) -> None:
    monkeypatch.delenv(
        "EDGE_MAX_XLSX_UNCOMPRESSED_MB",
        raising=False,
    )

    assert (
        get_max_xlsx_uncompressed_bytes()
        == 500 * 1024 * 1024
    )


def test_invalid_max_xlsx_uncompressed_size_configuration(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_XLSX_UNCOMPRESSED_MB",
        "invalid",
    )

    with pytest.raises(
        RuntimeError,
        match="positive integer",
    ):
        get_max_xlsx_uncompressed_bytes()


def test_default_max_xlsx_compression_ratio(
    monkeypatch,
) -> None:
    monkeypatch.delenv(
        "EDGE_MAX_XLSX_COMPRESSION_RATIO",
        raising=False,
    )

    assert (
        get_max_xlsx_compression_ratio()
        == 200
    )


def test_invalid_max_xlsx_compression_ratio_configuration(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_XLSX_COMPRESSION_RATIO",
        "invalid",
    )

    with pytest.raises(
        RuntimeError,
        match="positive integer",
    ):
        get_max_xlsx_compression_ratio()


def test_xlsx_exceeding_uncompressed_size_rejected(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_XLSX_UNCOMPRESSED_MB",
        "1",
    )

    monkeypatch.setenv(
        "EDGE_MAX_XLSX_COMPRESSION_RATIO",
        "1000000",
    )

    archive_buffer = BytesIO()

    with ZipFile(
        archive_buffer,
        mode="w",
        compression=ZIP_DEFLATED,
    ) as archive:
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            b"A" * (1024 * 1024 + 1),
        )

    xlsx_content = (
        archive_buffer.getvalue()
    )

    upload_file = create_upload_file(
        filename="oversized_archive.xlsx",
        content=xlsx_content,
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        asyncio.run(
            inspect_uploaded_file(
                upload_file
            )
        )

    assert (
        exception_info.value.status_code
        == 413
    )

    assert (
        exception_info.value.detail
        == (
            "The XLSX archive exceeds "
            "the configured maximum "
            "uncompressed size."
        )
    )


def test_xlsx_suspicious_compression_ratio_rejected(
    monkeypatch,
) -> None:
    monkeypatch.setenv(
        "EDGE_MAX_XLSX_COMPRESSION_RATIO",
        "10",
    )

    archive_buffer = BytesIO()

    with ZipFile(
        archive_buffer,
        mode="w",
        compression=ZIP_DEFLATED,
    ) as archive:
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            b"A" * 10000,
        )

    xlsx_content = (
        archive_buffer.getvalue()
    )

    upload_file = create_upload_file(
        filename="high_ratio.xlsx",
        content=xlsx_content,
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    with pytest.raises(
        HTTPException
    ) as exception_info:
        asyncio.run(
            inspect_uploaded_file(
                upload_file
            )
        )

    assert (
        exception_info.value.status_code
        == 413
    )

    assert (
        exception_info.value.detail
        == (
            "The XLSX archive contains "
            "a suspiciously compressed "
            "entry."
        )
    )